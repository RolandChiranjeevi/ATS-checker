"""DOCX extraction with python-docx."""

from __future__ import annotations

import unicodedata
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from cv_ats_check.extraction.errors import CorruptFile
from cv_ats_check.models import ExtractedDocument, PageInfo, TextBlock

_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _paragraph_to_block(paragraph, region: str = "body") -> TextBlock | None:
    text = unicodedata.normalize("NFC", paragraph.text.strip())
    if not text:
        return None
    style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
    runs = [r for r in paragraph.runs if r.text.strip()]
    bold = bool(runs) and all(r.bold for r in runs)
    size = next((r.font.size.pt for r in runs if r.font.size), None)
    return TextBlock(
        text=text,
        page=1,
        size=size,
        bold=bold,
        heading_style=style_name.startswith("Heading") or style_name == "Title",
        region=region,
    )


def _header_footer_paragraphs(part) -> list[Paragraph]:
    paragraphs = list(part.paragraphs)
    for table in part.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def _all_paragraphs(document) -> list[tuple[Paragraph, str]]:
    """Headers, then every body paragraph in document order, then footers (with their region).

    Walking the XML (instead of ``document.paragraphs``) also reaches text boxes and
    nested tables. ``mc:Fallback`` copies of text boxes are skipped to avoid duplicates.
    """
    headers: list[tuple[Paragraph, str]] = []
    footers: list[tuple[Paragraph, str]] = []
    for section in document.sections:
        if not section.header.is_linked_to_previous:
            headers += [(p, "header") for p in _header_footer_paragraphs(section.header)]
        if not section.footer.is_linked_to_previous:
            footers += [(p, "footer") for p in _header_footer_paragraphs(section.footer)]

    body = [
        (Paragraph(p, document), "body")
        for p in document.element.body.iter(qn("w:p"))
        if not any(ancestor.tag == _MC_FALLBACK for ancestor in p.iterancestors())
    ]
    return headers + body + footers


def extract_docx(path: Path) -> ExtractedDocument:
    """Extract text (incl. text boxes, headers and footers), tables and metadata from a DOCX."""
    try:
        document = Document(str(path))
    except Exception as exc:  # BadZipFile, PackageNotFoundError, KeyError, ...
        raise CorruptFile(f"Could not open DOCX (damaged or not a real DOCX): {path}") from exc

    blocks = [b for p, region in _all_paragraphs(document) if (b := _paragraph_to_block(p, region))]
    text = "\n".join(b.text for b in blocks)

    rels = document.part.rels.values()
    links = [r.target_ref for r in rels if r.reltype.endswith("/hyperlink") and r.is_external]
    images = sum(1 for r in rels if r.reltype.endswith("/image"))
    props = document.core_properties

    return ExtractedDocument(
        path=path,
        file_type="docx",
        text=text,
        pages=[PageInfo(number=1, char_count=len(text), image_count=images)],
        blocks=blocks,
        links=links,
        paginated=False,
        metadata={
            "title": (props.title or "").strip(),
            "author": (props.author or "").strip(),
            "last_modified_by": (props.last_modified_by or "").strip(),
            "created": props.created.date().isoformat() if props.created else "",
            "modified": props.modified.date().isoformat() if props.modified else "",
        },
        file_size=path.stat().st_size,
        table_count=sum(1 for _ in document.element.body.iter(qn("w:tbl"))),
    )
