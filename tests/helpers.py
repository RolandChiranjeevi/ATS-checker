"""Builders for synthetic test documents (no real personal data)."""

from __future__ import annotations

import io
import textwrap
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from cv_ats_check.models import ExtractedDocument, PageInfo, TextBlock

FIXTURES_DIR = Path(__file__).parent / "fixtures"

Line = tuple[str, str]  # (style, text); style is "name", "heading" or "body"
_FONTS = {
    "name": ("Helvetica-Bold", 18),
    "heading": ("Helvetica-Bold", 13),
    "body": ("Helvetica", 10),
}

EN_HEADINGS = ("Summary", "Skills", "Experience", "Education", "Projects")
SV_HEADINGS = ("Profil", "Kompetenser", "Arbetslivserfarenhet", "Utbildning", "Projekt")


def _new_canvas(path: Path, name: str | None, *, encrypt: bool = False) -> canvas.Canvas:
    protection = (
        StandardEncryption(userPassword="", ownerPassword="owner", canCopy=0) if encrypt else None
    )
    pdf = canvas.Canvas(str(path), pagesize=A4, encrypt=protection)
    if name:
        pdf.setTitle(f"{name} – CV")
        pdf.setAuthor(name)
    return pdf


def _photo_png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (200, 200), (120, 140, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def make_pdf(path: Path, lines: list[Line], *, photo: bool = False, encrypt: bool = False) -> Path:
    """Write a text PDF; headings are bold and larger so they look like headings."""
    name = next((text for style, text in lines if style == "name"), None)
    pdf = _new_canvas(path, name, encrypt=encrypt)
    _, height = A4
    if photo:  # about 3.4 % of the page
        pdf.drawImage(ImageReader(io.BytesIO(_photo_png())), 440, height - 160, 130, 130)
    y = height - 50
    for style, text in lines:
        font, size = _FONTS[style]
        wrapped = textwrap.wrap(text, width=95) or [""]
        if style == "heading":
            y -= 8
        for part in wrapped:
            if y < 60:
                pdf.showPage()
                y = height - 50
            pdf.setFont(font, size)
            pdf.drawString(50, y, part)
            y -= size * 1.45
    pdf.save()
    return path


def make_scanned_pdf(source_pdf: Path, target: Path) -> Path:
    """Rasterise the first page of ``source_pdf`` into an image-only PDF (no text layer)."""
    with pymupdf.open(source_pdf) as src:
        png = src[0].get_pixmap(dpi=80).tobytes("png")
    with pymupdf.open() as out:
        page = out.new_page()
        page.insert_image(page.rect, stream=png)
        out.save(target)
    return target


def make_two_column_pdf(path: Path) -> Path:
    """A sidebar layout: short skills lines on the left, experience lines on the right."""
    pdf = _new_canvas(path, "Anna Testsson")
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(40, 790, "Anna Testsson")
    left = [
        "Python Selenium Playwright Postman",
        "Jira and Confluence daily use",
        "Agile Scrum and Kanban teams",
        "ISTQB certified tester level",
        "Swedish native English fluent",
        "Stockholm based open to hybrid",
    ]
    right = [
        "QA Engineer at Example Software AB since 2022",
        "Designed and executed test plans for a customer portal",
        "Built automated end to end tests with Playwright",
        "Performed API testing with Postman in CI pipelines",
        "Reported and triaged defects together with developers",
        "Mentored two junior testers in exploratory testing",
    ]
    pdf.setFont("Helvetica", 9)
    for x, column in ((40, left), (250, right)):
        y = 740
        for text in column:
            pdf.drawString(x, y, text)
            y -= 16
    pdf.save()
    return path


def make_right_aligned_dates_pdf(path: Path) -> Path:
    """A single-column CV where short dates sit at the right margin (must not look 2-column)."""
    pdf = _new_canvas(path, "Anna Testsson")
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(50, 790, "Anna Testsson")
    jobs = [
        ("QA Engineer at Example Software AB in Stockholm", "2022–2025"),
        ("Test Analyst at Sample Systems AB in Gothenburg", "2019–2022"),
        ("Junior Tester at Demo Consulting AB in Malmö", "2016–2019"),
    ]
    bullets = [
        "- Designed and executed test plans for a customer portal with fifty thousand users.",
        "- Built automated end-to-end tests with Playwright and Python for every release.",
        "- Performed API testing with Postman and added the checks to a CI/CD pipeline.",
    ]
    y = 740
    for title, dates in jobs:
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(50, y, title)
        pdf.drawString(490, y, dates)
        y -= 16
        pdf.setFont("Helvetica", 10)
        for bullet in bullets:
            pdf.drawString(50, y, bullet)
            y -= 15
        y -= 10
    pdf.save()
    return path


def make_docx(
    path: Path,
    lines: list[Line],
    *,
    table: bool = False,
    header_text: str | None = None,
    columns: int = 1,
    columns_attribute: str | None = None,
) -> Path:
    """Write a DOCX where headings use real 'Heading 1' styles.

    ``columns`` sets the section to that many text columns; ``columns_attribute`` writes a raw
    (possibly malformed) ``w:num`` value instead.
    """
    document = Document()
    if columns > 1 or columns_attribute is not None:
        cols = document.sections[0]._sectPr.find(qn("w:cols"))
        cols.set(qn("w:num"), columns_attribute or str(columns))
    name = next((text for style, text in lines if style == "name"), None)
    if name:
        document.core_properties.author = name
        document.core_properties.title = f"{name} – CV"
    if header_text:
        document.sections[0].header.paragraphs[0].text = header_text
    if table:
        grid = document.add_table(rows=2, cols=2)
        for row, cells in enumerate((("Skills", "Python, Selenium"), ("Tools", "Jira, Postman"))):
            for col, value in enumerate(cells):
                grid.cell(row, col).text = value
    for style, text in lines:
        if style == "heading":
            document.add_heading(text, level=1)
        elif style == "name":
            document.add_heading(text, level=0)
        else:
            document.add_paragraph(text)
    document.save(str(path))
    return path


def cv_lines(
    *,
    lang: str = "en",
    contact: bool = True,
    dates: tuple[str, str, str] = ("2022–2025", "2019–2022", "2015–2018"),
) -> list[Line]:
    """A realistic synthetic CV. ``dates`` are used for job 1, job 2 and education."""
    d1, d2, d3 = dates
    if lang == "sv":
        name, headings = "Åsa Öberg", SV_HEADINGS
        city = "Göteborg"
        body = {
            "summary": "Testare med fem års erfarenhet av manuell och automatiserad testning av "
            "webbapplikationer. Van vid testplanering, utforskande testning och regressionstest.",
            "skills": "Python, Playwright, Selenium, API-testning, Postman, SQL, Jira, CI/CD, "
            "teststrategi, Scrum och agila metoder",
            "jobs": [
                (
                    f"QA-ingenjör, Exempel Mjukvara AB, Göteborg, {d1}",
                    [
                        "Tog fram testplaner och genomförde tester för en kundportal i Göteborg.",
                        "Byggde automatiserade end-to-end-tester med Playwright och Python.",
                        "Testade API:er med Postman och körde testerna i en CI/CD-pipeline.",
                        "Rapporterade och prioriterade fel i Jira tillsammans med utvecklare.",
                    ],
                ),
                (
                    f"Testanalytiker, Prov Systems AB, Västerås, {d2}",
                    [
                        "Genomförde manuella regressionstester inför varje release.",
                        "Skrev SQL-frågor för att verifiera dataintegritet mellan tjänster.",
                        "Dokumenterade testfall och spårbarhet mellan krav och testresultat.",
                    ],
                ),
            ],
            "education": [f"Civilingenjör datateknik, Exempeluniversitetet, Göteborg, {d3}"],
            "projects": [
                "CV-kontroll: lokalt Python-verktyg som analyserar CV:n och rapporterar ATS-problem."
            ],
        }
    else:
        name, headings = "Anna Testsson", EN_HEADINGS
        city = "Stockholm"
        body = {
            "summary": "QA engineer with five years of experience in manual and automated testing "
            "of web applications. Skilled in test planning, exploratory testing and building "
            "reliable automated regression suites.",
            "skills": "Python, Playwright, Selenium, API testing, Postman, SQL, Jira, CI/CD, "
            "GitHub Actions, test strategy, regression testing, Agile and Scrum",
            "jobs": [
                (
                    f"QA Engineer, Example Software AB, Stockholm, {d1}",
                    [
                        "Designed and executed test plans for a customer portal with fifty thousand users.",
                        "Built automated end-to-end tests with Playwright and Python, cutting regression time by sixty percent.",
                        "Performed API testing with Postman and added the checks to a CI/CD pipeline on GitHub Actions.",
                        "Reported and triaged defects in Jira with developers and product owners in an Agile team.",
                        "Mentored two junior testers in exploratory testing and test case design.",
                    ],
                ),
                (
                    f"Test Analyst, Sample Systems AB, Gothenburg, {d2}",
                    [
                        "Executed manual regression and acceptance tests for a payment platform before every release.",
                        "Wrote SQL queries to verify data integrity across several back end services.",
                        "Documented test cases and kept traceability between requirements and test results.",
                        "Worked with customer support to reproduce reported issues and verify the fixes.",
                    ],
                ),
            ],
            "education": [
                f"BSc in Computer Science, Example University, Stockholm, {d3}",
                "ISTQB Certified Tester Foundation Level, 2021",
            ],
            "projects": [
                "CV checker: a local Python tool that analyses CVs and reports ATS issues.",
                "Playwright demo suite: end-to-end tests for a sample web shop run on GitHub Actions.",
            ],
        }
    summary_h, skills_h, exp_h, edu_h, proj_h = headings

    lines: list[Line] = [("name", name)]
    if contact:
        lines.append(
            ("body", f"{name.split()[0].lower()}.example@example.com | +46 70 123 45 67 | {city}")
        )
        lines.append(("body", "linkedin.com/in/exampleperson | github.com/exampleperson"))
    lines += [("heading", summary_h), ("body", body["summary"])]
    lines += [("heading", skills_h), ("body", body["skills"])]
    lines.append(("heading", exp_h))
    for title, bullets in body["jobs"]:
        lines.append(("body", title))
        lines += [("body", f"- {b}") for b in bullets]
    lines.append(("heading", edu_h))
    lines += [("body", e) for e in body["education"]]
    lines += [("heading", proj_h)] + [("body", f"- {p}") for p in body["projects"]]
    return lines


def make_doc(lines: list[str] | list[TextBlock], *, bold_all: bool = False) -> ExtractedDocument:
    """Build an in-memory document for unit tests of checks (no file needed)."""
    blocks = [
        line if isinstance(line, TextBlock) else TextBlock(line, 1, size=10.0, bold=bold_all)
        for line in lines
    ]
    text = "\n".join(b.text for b in blocks)
    return ExtractedDocument(
        path=Path("memory.pdf"),
        file_type="pdf",
        text=text,
        pages=[PageInfo(1, len(text), 0)],
        blocks=blocks,
    )


def build_all(directory: Path = FIXTURES_DIR) -> None:
    """Create every binary fixture that is missing (idempotent)."""
    directory.mkdir(parents=True, exist_ok=True)

    def missing(name: str) -> Path | None:
        target = directory / name
        return None if target.exists() else target

    if target := missing("cv_normal_en.pdf"):
        make_pdf(target, cv_lines())
    if target := missing("cv_no_contact.pdf"):
        make_pdf(target, cv_lines(contact=False))
    if target := missing("cv_mixed_dates.pdf"):
        make_pdf(target, cv_lines(dates=("Jan 2022 - Mar 2025", "2019/2021", "2015–2018")))
    if target := missing("cv_swedish.pdf"):
        make_pdf(target, cv_lines(lang="sv"))
    if target := missing("cv_scanned.pdf"):
        make_scanned_pdf(directory / "cv_normal_en.pdf", target)
    if target := missing("cv_normal.docx"):
        make_docx(target, cv_lines())
    if target := missing("cv_two_column.pdf"):
        make_two_column_pdf(target)
    if target := missing("cv_right_aligned_dates.pdf"):
        make_right_aligned_dates_pdf(target)
    if target := missing("cv_photo.pdf"):
        make_pdf(target, cv_lines(), photo=True)
    if target := missing("cv_owner_protected.pdf"):
        make_pdf(target, cv_lines(), encrypt=True)
    if target := missing("cv_two_column.docx"):
        make_docx(target, cv_lines(), columns=2)
    if target := missing("cv_table.docx"):
        make_docx(target, cv_lines(), table=True)
    if target := missing("cv_header_contact.docx"):
        make_docx(
            target,
            cv_lines(contact=False),
            header_text="anna.example@example.com | +46 70 123 45 67",
        )
    if target := missing("cv_corrupt.pdf"):
        target.write_bytes(b"%PDF-1.4\nthis is definitely not a real pdf\n")
    if target := missing("cv_empty.pdf"):
        target.write_bytes(b"")
