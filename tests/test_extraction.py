"""Extraction: positive and negative file cases."""

from __future__ import annotations

import pytest

from cv_ats_check.extraction import (
    CorruptFile,
    EmptyFile,
    ExtractionError,
    MissingFile,
    UnsupportedFileType,
    extract_document,
)


def test_normal_pdf_has_text_pages_and_styled_blocks(fixtures_dir):
    doc = extract_document(fixtures_dir / "cv_normal_en.pdf")
    assert doc.file_type == "pdf"
    assert "Anna Testsson" in doc.text
    assert doc.page_count >= 1
    assert doc.total_chars > 500
    assert any(b.bold and b.text == "Experience" for b in doc.blocks)


def test_swedish_characters_survive_extraction(fixtures_dir):
    doc = extract_document(fixtures_dir / "cv_swedish.pdf")
    for expected in ("Åsa Öberg", "Göteborg", "Arbetslivserfarenhet", "Västerås", "Kompetenser"):
        assert expected in doc.text


def test_scanned_pdf_has_images_but_no_text(fixtures_dir):
    doc = extract_document(fixtures_dir / "cv_scanned.pdf")
    assert doc.image_count >= 1
    assert doc.total_chars < 20


def test_docx_extraction_marks_heading_styles(fixtures_dir):
    doc = extract_document(fixtures_dir / "cv_normal.docx")
    assert doc.file_type == "docx"
    assert doc.paginated is False
    assert any(b.heading_style and b.text == "Experience" for b in doc.blocks)


_TEXTBOX_XML = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps"><w:pict><v:textbox><w:txbxContent>
      <w:p><w:r><w:t>TextboxOnlyWord</w:t></w:r></w:p>
    </w:txbxContent></v:textbox></w:pict></mc:Choice>
    <mc:Fallback><w:pict><v:textbox><w:txbxContent>
      <w:p><w:r><w:t>TextboxOnlyWord</w:t></w:r></w:p>
    </w:txbxContent></v:textbox></w:pict></mc:Fallback>
  </mc:AlternateContent>
</w:r>
"""


def _docx_with(tmp_path, customise):
    from docx import Document

    document = Document()
    document.add_paragraph("Body paragraph")
    customise(document)
    path = tmp_path / "cv.docx"
    document.save(str(path))
    return path


def test_docx_text_box_text_is_extracted_once(tmp_path):
    from docx.oxml import parse_xml

    path = _docx_with(tmp_path, lambda d: d.paragraphs[0]._p.append(parse_xml(_TEXTBOX_XML)))
    doc = extract_document(path)
    assert doc.text.count("TextboxOnlyWord") == 1


def test_docx_header_and_footer_text_is_extracted(tmp_path):
    def customise(document):
        document.sections[0].header.paragraphs[0].text = "Header Person +46 70 123 45 67"
        document.sections[0].footer.paragraphs[0].text = "Footer note"

    doc = extract_document(_docx_with(tmp_path, customise))
    lines = doc.lines
    assert lines[0] == "Header Person +46 70 123 45 67"
    assert lines[-1] == "Footer note"
    assert "Body paragraph" in lines


def test_docx_nested_table_text_is_extracted(tmp_path):
    def customise(document):
        outer = document.add_table(rows=1, cols=1)
        outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "NestedCellWord"

    assert "NestedCellWord" in extract_document(_docx_with(tmp_path, customise)).text


def test_empty_file_is_rejected(fixtures_dir):
    with pytest.raises(EmptyFile):
        extract_document(fixtures_dir / "cv_empty.pdf")


def test_corrupt_pdf_is_rejected(fixtures_dir):
    with pytest.raises(CorruptFile):
        extract_document(fixtures_dir / "cv_corrupt.pdf")


def test_corrupt_docx_is_rejected(tmp_path):
    bad = tmp_path / "cv.docx"
    bad.write_bytes(b"not a zip file")
    with pytest.raises(CorruptFile):
        extract_document(bad)


@pytest.mark.parametrize("name", ["cv.txt", "cv.png", "cv.doc", "cv", "cv.pdf.exe"])
def test_wrong_file_type_is_rejected(tmp_path, name):
    path = tmp_path / name
    path.write_bytes(b"some content")
    with pytest.raises(UnsupportedFileType):
        extract_document(path)


def test_missing_file_is_rejected(tmp_path):
    with pytest.raises(MissingFile):
        extract_document(tmp_path / "nope.pdf")


def test_all_errors_share_a_base_class():
    for cls in (MissingFile, UnsupportedFileType, EmptyFile, CorruptFile):
        assert issubclass(cls, ExtractionError)


def test_uppercase_extension_is_accepted(fixtures_dir, tmp_path):
    copy = tmp_path / "CV.PDF"
    copy.write_bytes((fixtures_dir / "cv_normal_en.pdf").read_bytes())
    assert "Anna Testsson" in extract_document(copy).text
